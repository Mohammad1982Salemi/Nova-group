"""One review workbook per language: a sheet for the site and a sheet for the reports.

    python tools/review/workbook_all.py <lang> <out.xlsx>

Reads the extracted tables and the fix lists this review produced, and lays every unit beside
its Persian, English and reviewed text, marking the rows that were changed and saying why.
The point is that someone who knows the language can confirm the work in one pass without
opening the site, the reports or the code.
"""
import json
import re
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[2]
REVIEW = ROOT / "_originals" / "review"
HEAD = ["#", "جای دقیق", "فارسی", "English", "زبان هدف", "وضعیت", "یادداشت"]
WIDTH = [5, 34, 52, 52, 52, 14, 62]
FIXED = PatternFill("solid", fgColor="FFF6DA")


def norm(text):
    return re.sub(r"\s+", " ", text or "").strip()


def fix_map(paths):
    """Every fix this review applied, under both the key and the position it can be found by."""
    out = {}
    for path in paths:
        if not path.exists():
            continue
        for fix in json.loads(path.read_text(encoding="utf-8")):
            why = fix["why"]
            if "en" in fix:
                out[norm(fix["en"])] = why
            if "index" in fix:
                out[f"{fix['section']}[{fix['index']}]"] = why
    return out


def sheet(book, title, units, lang, fixes):
    page = book.create_sheet(title)
    page.sheet_view.rightToLeft = True
    for col, (name, width) in enumerate(zip(HEAD, WIDTH), start=1):
        cell = page.cell(row=1, column=col, value=name)
        cell.font = Font(name="Arial", bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F3864")
        cell.alignment = Alignment(horizontal="center", vertical="center")
        page.column_dimensions[get_column_letter(col)].width = width
    page.row_dimensions[1].height = 22

    changed = 0
    for number, unit in enumerate(units, start=1):
        place = f"{unit['kind']}[{unit['index']}]" if "index" in unit else None
        why = fixes.get(norm(unit.get("en"))) or (fixes.get(place) if place else None)
        fixed = bool(why)
        changed += fixed
        where = (unit.get("where") or [f"{unit.get('src')}:{unit.get('line')}"])[0]
        row = [number, where, unit.get("fa") or "", unit.get("en") or "", unit.get(lang) or "",
               "اصلاح شد" if fixed else "کنترل شد", why or ""]
        for col, value in enumerate(row, start=1):
            cell = page.cell(row=number + 1, column=col, value=value)
            cell.font = Font(name="Tahoma" if col == 3 else "Arial", size=10)
            cell.alignment = Alignment(vertical="top", wrap_text=True,
                                       horizontal="center" if col in (1, 6) else "left")
            if fixed:
                cell.fill = FIXED
    page.freeze_panes = "A2"
    page.auto_filter.ref = f"A1:{get_column_letter(len(HEAD))}{len(units) + 1}"
    return changed


def main():
    lang, out = sys.argv[1], sys.argv[2]
    site = json.loads((REVIEW / "units.json").read_text(encoding="utf-8"))
    reports = json.loads((REVIEW / "report-units.json").read_text(encoding="utf-8"))

    site_fixes = fix_map([REVIEW / f"{lang}_site_fixes.json"])
    report_fixes = fix_map(sorted(REVIEW.glob(f"{lang}_cme_fixes*.json")))

    book = Workbook()
    book.remove(book.active)
    site_changed = sheet(book, f"سایت {lang}", site, lang, site_fixes)
    report_changed = sheet(book, f"گزارش‌ها {lang}", reports, lang, report_fixes)

    note = book.create_sheet("راهنما")
    note.sheet_view.rightToLeft = True
    lines = [
        ("راهنمای این جدول", True), ("", False),
        (f"برگ «سایت» همهٔ {len(site)} واحد ترجمهٔ سایت را کنار هم می‌گذارد و برگ «گزارش‌ها» هر {len(reports)} واحد", False),
        ("واژه‌نامهٔ مشترک چهار گزارش CME و متن گزارش صندوق طلا را. ستون‌ها: جای دقیق در متن، متن پارسی", False),
        ("(زبان مرجع)، متن انگلیسی، متن زبان هدف، وضعیت و دلیل اصلاح.", False),
        (f"سطرهای زردرنگ اصلاح شده‌اند: {site_changed} سطر در سایت و {report_changed} سطر در گزارش‌ها.", False),
        ("", False),
        ("برای کسی که زبان می‌داند: کافی است ستون «زبان هدف» را با ستون انگلیسی بسنجد؛ اگر جایی معنا یکی نبود", False),
        ("یا اصطلاح غلط بود، همان سطر را علامت بزند.", False),
        ("", False),
        ("تاریخ: ۲۰ مهر ۱۴۰۵ / 2026-10-12", False),
    ]
    for i, (text, bold) in enumerate(lines, start=1):
        cell = note.cell(row=i, column=1, value=text)
        cell.font = Font(name="Tahoma" if bold else "Arial", size=12 if bold else 10, bold=bold)
        cell.alignment = Alignment(horizontal="right")
    note.column_dimensions["A"].width = 112

    book.save(out)
    print(f"{out}: site {len(site)} units ({site_changed} revised), reports {len(reports)} units ({report_changed} revised)")


if __name__ == "__main__":
    main()
